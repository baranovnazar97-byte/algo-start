"""Собирает пояснительную записку Word из документации проекта."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
OUTPUT = DOCS / "Курсовой_проект_Алгоритмика.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(0)

    for index in range(1, 4):
        style = styles[f"Heading {index}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.color.rgb = RGBColor(31, 41, 55)
        style.font.bold = True
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 2"].font.size = Pt(15)
    styles["Heading 3"].font.size = Pt(14)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(14)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.left_indent = Cm(1.25)
        style.paragraph_format.first_line_indent = Cm(0)

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code_style.font.size = Pt(9)
    code_style.paragraph_format.left_indent = Cm(0.5)
    code_style.paragraph_format.right_indent = Cm(0.5)
    code_style.paragraph_format.first_line_indent = Cm(0)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(3)
    code_style.paragraph_format.line_spacing = 1.0


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE")


def add_title_page(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("[НАЗВАНИЕ ОБРАЗОВАТЕЛЬНОЙ ОРГАНИЗАЦИИ]\n")
    run.bold = True
    run.font.size = Pt(14)
    p.add_run("[ОТДЕЛЕНИЕ / ЦИКЛОВАЯ КОМИССИЯ]").font.size = Pt(12)

    document.add_paragraph("\n\n")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("КУРСОВОЙ ПРОЕКТ")
    run.bold = True
    run.font.size = Pt(18)
    p.add_run("\nпо дисциплине [НАЗВАНИЕ ДИСЦИПЛИНЫ]").font.size = Pt(14)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run("ООО «Алгоритмика» (СБЕР)\n")
    run.bold = True
    run.font.size = Pt(17)
    run = p.add_run("«Игровой тренажёр по алгоритмам.\nМини-игры для обучения детей алгоритмическому мышлению»")
    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph("\n")
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table.autofit = False
    table.columns[0].width = Cm(5)
    table.columns[1].width = Cm(8)
    left, right = table.rows[0].cells
    left.text = ""
    paragraph = right.paragraphs[0]
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(
        "Выполнили:\n"
        "[ФИО участника 1], группа [___]\n"
        "[ФИО участника 2], группа [___]\n"
        "[ФИО участника 3], группа [___]\n\n"
        "Руководитель:\n"
        "[ФИО, должность]"
    )
    # Убираем внешние границы служебной таблицы.
    for cell in table.rows[0].cells:
        properties = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), "nil")
            borders.append(node)
        properties.append(borders)

    document.add_paragraph("\n\n")
    p = document.add_paragraph("[ГОРОД] — 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(24)
    document.add_page_break()


def add_assignment_page(document: Document) -> None:
    document.add_heading("Задание на курсовой проект", level=1)
    lines = [
        ("Тема", "Игровой тренажёр по алгоритмам для детей"),
        ("Цель", "Разработать многостраничное веб-приложение с четырьмя простыми мини-играми, учётом очков и прогресса."),
        ("Технологии", "React, TypeScript, NestJS, PostgreSQL, pg, Docker Compose"),
        ("Исходные требования", "Регистрация ребёнка; выдача игровых заданий; прохождение уровней; начисление очков; открытие новых заданий; страница прогресса; детский минималистичный интерфейс; документация Word."),
        ("Состав результата", "Исходный код frontend/backend, SQL-инициализация, Docker-конфигурация, инструкции по развёртыванию и обновлению, тест-план, руководство пользователя, пояснительная записка."),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "Пункт"
    table.rows[0].cells[1].text = "Содержание"
    set_repeat_table_header(table.rows[0])
    for title, content in lines:
        cells = table.add_row().cells
        cells[0].text = title
        cells[1].text = content
    format_table(table)
    document.add_paragraph()
    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Дата выдачи: ").bold = True
    p.add_run("[ДД.ММ.2026]    ")
    p.add_run("Срок сдачи: ").bold = True
    p.add_run("[ДД.ММ.2026]")
    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Подпись руководителя: ____________    Подписи студентов: ____________")
    document.add_page_break()


def add_toc(document: Document) -> None:
    document.add_heading("Содержание", level=1)
    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Щёлкните правой кнопкой и выберите «Обновить поле»")
    document.add_paragraph(
        "Примечание: после открытия в Microsoft Word нажмите Ctrl+A, затем F9, чтобы обновить содержание и номера страниц."
    )
    document.add_page_break()


INLINE_TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*|\[[^]]+\]\([^)]+\)|https?://\S+)")


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(11)
        elif token.startswith("["):
            label, url = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            paragraph.add_run(f"{label} ({url})")
        else:
            paragraph.add_run(token.rstrip(".,;"))
            suffix = token[len(token.rstrip(".,;")) :]
            if suffix:
                paragraph.add_run(suffix)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def format_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(10)
                    if row_index == 0:
                        run.bold = True


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        add_inline(table.rows[0].cells[index].paragraphs[0], value.strip())
    set_repeat_table_header(table.rows[0])
    for source_row in rows[1:]:
        cells = table.add_row().cells
        for index in range(width):
            add_inline(cells[index].paragraphs[0], source_row[index].strip() if index < len(source_row) else "")
    format_table(table)


def markdown_to_document(document: Document, path: Path, skip_first_heading: bool = False) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    first_heading_skipped = False
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            p = document.add_paragraph(style="Code Block")
            if language:
                r = p.add_run(f"[{language}]\n")
                r.bold = True
                r.font.color.rgb = RGBColor(66, 85, 99)
            p.add_run("\n".join(code_lines))
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[index + 1]):
            table_lines = [stripped]
            index += 2  # пропускаем строку-разделитель Markdown
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
            add_table(document, rows)
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            title = heading.group(2).strip()
            if skip_first_heading and not first_heading_skipped:
                first_heading_skipped = True
            else:
                document.add_heading(title, level=level)
            index += 1
            continue

        if stripped.startswith("> "):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.right_indent = Cm(1)
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, stripped[2:])
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        number = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or number:
            p = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(p, (bullet or number).group(1))
            index += 1
            continue

        paragraph_parts = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", "```", "> ", "|"))
                or re.match(r"^[-*]\s+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_parts.append(candidate)
            index += 1
        p = document.add_paragraph()
        add_inline(p, " ".join(paragraph_parts))


def add_team_signature_page(document: Document) -> None:
    document.add_page_break()
    document.add_heading("Лист распределения ответственности", level=1)
    document.add_paragraph(
        "Страница заполняется фактическими данными перед сдачей. Указанные роли — рекомендуемое распределение зон, а не автоматически подтверждённое авторство конкретных файлов."
    )
    rows = [
        ["Участник", "ФИО", "Зона ответственности", "Фактически выполнено", "Подпись"],
        ["1", "[ФИО]", "Тестирование, deploy и документация", "[Заполнить]", ""],
        ["2", "[ФИО]", "Frontend и дизайн", "[Заполнить]", ""],
        ["3", "[ФИО]", "Backend и база данных", "[Заполнить]", ""],
    ]
    add_table(document, rows)


def build() -> None:
    document = Document()
    configure_document(document)
    add_title_page(document)
    add_assignment_page(document)
    add_toc(document)

    markdown_to_document(document, DOCS / "COURSE_PROJECT.md", skip_first_heading=True)

    document.add_page_break()
    markdown_to_document(document, DOCS / "USER_GUIDE.md")

    document.add_page_break()
    markdown_to_document(document, DOCS / "TEST_PLAN.md")

    document.add_page_break()
    markdown_to_document(document, ROOT / "WEEKLY_PLAN.md")

    add_team_signature_page(document)

    for section in document.sections:
        add_page_number(section)

    properties = document.core_properties
    properties.title = "Курсовой проект — игровой тренажёр по алгоритмам"
    properties.subject = "React, TypeScript, NestJS, PostgreSQL, Docker"
    properties.author = "Участники курсового проекта — заполнить перед сдачей"
    properties.keywords = "алгоритмы, обучение, мини-игры, React, NestJS, PostgreSQL"

    document.save(OUTPUT)
    # ASCII-only output keeps the script usable in the default Windows console,
    # whose legacy code page may not be able to print a Cyrillic file name.
    print("DOCX created successfully")


if __name__ == "__main__":
    build()
